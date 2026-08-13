"""Upload extraction is a whitelist: recognised fields in, everything else out."""

import pytest

from medication_safety.analysis import assess_profile
from medication_safety.extraction import (
    extract_from_csv,
    extract_from_json,
    extract_from_text,
    extract_upload,
)

CLINIC_NOTE = """
Patient name: Ahmed Al-Otaibi
MRN 4482910
Date of birth: 04/11/1947

Age: 78
INR 1.6, eGFR 42, serum creatinine 1.6
Heart rate 58

Current medications
- Warfarin 5 mg once daily
- Amiodarone 200 mg once daily
- Ketoconazole 200 mg once daily oral
- Simvastatin 40 mg at bedtime
- Zopiclone 7.5 mg at bedtime
"""


def test_recognised_medications_and_parameters_are_filled() -> None:
    result = extract_from_text(CLINIC_NOTE)
    names = {m["name"] for m in result.profile["medications"]}

    assert {"warfarin", "amiodarone", "ketoconazole", "simvastatin"} <= names
    assert result.profile["age"] == 78
    assert result.profile["inr"] == 1.6
    assert result.profile["egfr"] == 42
    assert result.profile["heart_rate"] == 58


def test_identifiers_are_never_carried_into_the_profile() -> None:
    result = extract_from_text(CLINIC_NOTE)
    blob = str(result.profile).lower()

    for identifier in ("ahmed", "al-otaibi", "4482910", "1947", "mrn"):
        assert identifier not in blob


def test_extracted_profile_passes_the_identifier_guardrail() -> None:
    # The point of the whitelist: a note full of identifiers still yields a
    # profile that analysis accepts.
    result = extract_from_text(CLINIC_NOTE)
    assessment = assess_profile(result.profile)
    assert assessment.findings


def test_dose_and_frequency_are_captured() -> None:
    result = extract_from_text(CLINIC_NOTE)
    warfarin = next(m for m in result.profile["medications"] if m["name"] == "warfarin")
    assert warfarin["dose"] == "5 mg"
    assert warfarin["frequency"] == "once daily"

    keto = next(m for m in result.profile["medications"] if m["name"] == "ketoconazole")
    assert keto["route"] == "oral"


def test_unknown_medicine_is_reported_not_silently_dropped() -> None:
    result = extract_from_text(CLINIC_NOTE)
    assert any("Zopiclone" in line for line in result.unrecognised)
    assert any("not in the recognised vocabulary" in note for note in result.notes)


def test_discarded_lines_are_counted() -> None:
    result = extract_from_text(CLINIC_NOTE)
    assert result.ignored_lines > 0
    assert any("never leave the server" in note for note in result.notes)


def test_longest_medication_name_wins() -> None:
    result = extract_from_text("Metoprolol succinate 50 mg once daily")
    assert result.profile["medications"][0]["name"] == "metoprolol succinate"


def test_bare_numbers_are_not_adopted_as_values() -> None:
    result = extract_from_text("The ward has 42 beds and 78 staff.\nWarfarin 5 mg")
    assert "age" not in result.profile
    assert "egfr" not in result.profile


def test_negated_dialysis_is_read_as_false() -> None:
    assert extract_from_text("Patient is not on dialysis").profile["dialysis"] is False
    assert extract_from_text("Receiving dialysis twice weekly").profile["dialysis"] is True


def test_csv_keeps_named_rows_and_flags_unknown_ones() -> None:
    result = extract_from_csv(
        "name,dose,frequency\nWarfarin,5 mg,once daily\nZopiclone,7.5 mg,at bedtime\n"
    )
    assert len(result.profile["medications"]) == 2
    assert result.unrecognised == ["Zopiclone"]


def test_json_drops_unsupported_fields() -> None:
    result = extract_from_json(
        '{"age": 70, "patient_name": "X", "mrn": "1", '
        '"medications": [{"name": "Warfarin"}]}'
    )
    assert result.profile["age"] == 70
    assert "patient_name" not in result.profile
    assert "mrn" not in result.profile
    assert any("dropped" in note for note in result.notes)


def test_invalid_json_is_rejected_clearly() -> None:
    with pytest.raises(ValueError, match="Invalid JSON"):
        extract_from_json("{nope")


def test_unsupported_file_type_is_rejected() -> None:
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_upload("book.xlsx", "application/vnd.ms-excel", b"PK\x03\x04")


def test_oversized_and_empty_uploads_are_rejected() -> None:
    with pytest.raises(ValueError, match="2 MB"):
        extract_upload("big.txt", "text/plain", b"x" * (2 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="empty"):
        extract_upload("empty.txt", "text/plain", b"")


def test_image_without_a_vision_model_says_what_to_do(monkeypatch) -> None:
    monkeypatch.delenv("OPENROUTER_VISION_MODEL", raising=False)
    with pytest.raises(ValueError, match="OPENROUTER_VISION_MODEL"):
        extract_upload("list.png", "image/png", b"\x89PNG\r\n\x1a\n fake")


def test_upload_dispatches_on_suffix() -> None:
    assert extract_upload("a.csv", "text/csv", b"name\nWarfarin\n").source == "csv"
    assert extract_upload("a.json", "application/json", b"{}").source == "json"
    assert extract_upload("a.md", "text/markdown", b"Warfarin 5 mg").source == "markdown"


# ---------- PDF and Word ----------

MED_LINES = [
    "Age: 78",
    "INR 1.6",
    "Warfarin 5 mg once daily",
    "Ketoconazole 200 mg once daily oral",
    "Simvastatin 40 mg at bedtime",
]


def _build_pdf(lines: list[str]) -> bytes:
    """A minimal real PDF with a text layer, written with pypdf."""
    import io

    from pypdf import PdfWriter
    from pypdf.generic import (
        ArrayObject,
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
        NumberObject,
    )

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)

    stream = "BT /F1 12 Tf 72 720 Td 16 TL\n"
    for line in lines:
        stream += f"({line}) Tj T*\n"
    stream += "ET"

    content = DecodedStreamObject()
    content.set_data(stream.encode("latin-1"))
    page[NameObject("/Contents")] = writer._add_object(content)

    font = DictionaryObject()
    font.update(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    resources = DictionaryObject()
    fonts = DictionaryObject()
    fonts[NameObject("/F1")] = writer._add_object(font)
    resources[NameObject("/Font")] = fonts
    page[NameObject("/Resources")] = resources
    page[NameObject("/MediaBox")] = ArrayObject(
        [NumberObject(0), NumberObject(0), NumberObject(612), NumberObject(792)]
    )

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _build_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    import io

    import docx

    document = docx.Document()
    for line in paragraphs:
        document.add_paragraph(line)
    if table_rows:
        table = document.add_table(rows=0, cols=len(table_rows[0]))
        for row in table_rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_text_layer_is_read() -> None:
    result = extract_upload("list.pdf", "application/pdf", _build_pdf(MED_LINES))
    names = {m["name"] for m in result.profile["medications"]}

    assert result.source == "pdf"
    assert {"warfarin", "ketoconazole", "simvastatin"} <= names
    assert result.profile["age"] == 78
    assert any("text layer of 1 page" in note for note in result.notes)


def test_scanned_pdf_without_text_is_reported_clearly() -> None:
    blank = _build_pdf([])
    with pytest.raises(ValueError, match="no text layer"):
        extract_upload("scan.pdf", "application/pdf", blank)


def test_corrupt_pdf_is_rejected_not_crashed() -> None:
    with pytest.raises(ValueError, match="Could not read this PDF"):
        extract_upload("broken.pdf", "application/pdf", b"%PDF-1.4 truncated")


def test_word_paragraphs_are_read() -> None:
    data = _build_docx(MED_LINES)
    result = extract_upload("list.docx", "", data)
    names = {m["name"] for m in result.profile["medications"]}

    assert result.source == "word"
    assert {"warfarin", "ketoconazole", "simvastatin"} <= names
    assert result.profile["inr"] == 1.6


def test_word_tables_are_read_too() -> None:
    # A medication list is very often a table, and paragraph text alone
    # would miss every row.
    data = _build_docx(
        ["Medication review"],
        [
            ["Medicine", "Dose", "Frequency"],
            ["Warfarin", "5 mg", "once daily"],
            ["Amiodarone", "200 mg", "once daily"],
        ],
    )
    result = extract_upload("table.docx", "", data)
    names = {m["name"] for m in result.profile["medications"]}

    assert {"warfarin", "amiodarone"} <= names
    assert any("table(s)" in note for note in result.notes)


def test_legacy_doc_format_says_what_to_do() -> None:
    with pytest.raises(ValueError, match=r"\.docx or PDF"):
        extract_upload("old.doc", "application/msword", b"\xd0\xcf\x11\xe0")


def test_identifiers_in_a_pdf_are_still_discarded() -> None:
    data = _build_pdf(["Patient name: Ahmed", "MRN 4482910"] + MED_LINES)
    result = extract_upload("note.pdf", "application/pdf", data)
    blob = str(result.profile).lower()
    assert "ahmed" not in blob
    assert "4482910" not in blob
